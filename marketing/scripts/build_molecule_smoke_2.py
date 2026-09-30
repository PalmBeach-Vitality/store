#!/usr/bin/env python3
"""Build the 2-row molecule smoke-test sheet `23-molecule-smoke-2`.

Row 1 stills on GPT Image 2.5 Sunburst (OpenRouter), row 2 on Grok Imagine Image 2.0.
Both rows fill every still column, so switching engines is one `model_still` edit.
Video is fal Kling v3 Pro, 15s + 15s, joined by Creatomate. Audio stays off: the
prep nodes add SILENT_PREFIX and send generate_audio=false, so the sheet omits both.

Does NOT modify 13-chem-breakdown-54.csv.

Output:
  marketing/sheets/23-molecule-smoke-2.csv
  --xlsx PATH  also writes the upload copy (every string cell text-formatted, so
               Google Sheets keeps `9:16` as text instead of a time)
"""

from __future__ import annotations

import argparse
import csv
import re
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "sheets" / "23-molecule-smoke-2.csv"
TAB = "23-molecule-smoke-2"

FIELDS = [
    "creation_id",
    "rank",
    "status",
    "compound_name",
    "look_name",
    "model_still",
    "still_prompt",
    "aspect_ratio",
    "still_resolution",
    "still_size",
    "still_quality",
    "still_n",
    "still_timeout_seconds",
    "model_video",
    "resolution",
    "hop1_duration_seconds",
    "video_motion_prompt",
    "hop2_duration_seconds",
    "extend_motion_prompt",
    "negative_prompt",
    "cfg_scale",
    "video_poll_seconds",
    "video_max_wait_seconds",
    "render_width",
    "render_height",
    "render_frame_rate",
    "video_format",
    "last_frame_format",
    "creatomate_poll_seconds",
    "creatomate_max_polls",
    "times_used",
    "last_used_at",
    "still_url",
    "hop1_video_url",
    "last_frame_url",
    "hop2_video_url",
    "video_url",
]

NUMERIC = {
    "rank",
    "still_n",
    "still_timeout_seconds",
    "hop1_duration_seconds",
    "hop2_duration_seconds",
    "cfg_scale",
    "video_poll_seconds",
    "video_max_wait_seconds",
    "render_width",
    "render_height",
    "render_frame_rate",
    "creatomate_poll_seconds",
    "creatomate_max_polls",
    "times_used",
}

OUTPUT_COLUMNS = {
    "last_used_at",
    "still_url",
    "hop1_video_url",
    "last_frame_url",
    "hop2_video_url",
    "video_url",
}

PROMPT_COLUMNS = ("still_prompt", "video_motion_prompt", "extend_motion_prompt", "negative_prompt")

SILENT_PREFIX = "Silent video. No soundtrack, no music, no sound effects, no dialogue, no ambient audio."
KLING_PROMPT_MAX = 2500

STILL_HEADER = (
    "Photorealistic cinematic science-fiction film still, vertical 9:16 frame. "
    "Hyperreal, as if a real cinema camera with a macro probe lens were floating inside "
    "living tissue: physically accurate light, refraction, reflections and depth of field. "
    "Premium feature-film visual effects, like the hero shot of a big-budget science-fiction movie. "
)

STILL_LOCK = (
    "The image contains no text, letters, numbers, labels, logos or watermarks, "
    "and no people, hands, vials or pens."
)

NEGATIVE = (
    "text, letters, numbers, labels, captions, logos, watermark, people, hands, faces, "
    "vials, pens, bottles, cartoon, clay, plastic toy look, flat diagram, infographic, "
    "grainy microscope footage, low resolution, blur, noise, flicker, morphing, warping, "
    "distortion, jitter, cuts, scene change, fade to black, low quality"
)

SHARED = {
    "status": "Active",
    "aspect_ratio": "9:16",
    "still_resolution": "2k",
    "still_size": "1440x2560",
    "still_quality": "high",
    "still_n": 1,
    "still_timeout_seconds": 600,
    "model_video": "fal-ai/kling-video/v3/pro/image-to-video",
    "resolution": "1080p",
    "hop1_duration_seconds": 15,
    "hop2_duration_seconds": 15,
    "negative_prompt": NEGATIVE,
    "cfg_scale": 0.5,
    "video_poll_seconds": 5,
    "video_max_wait_seconds": 900,
    "render_width": 1080,
    "render_height": 1920,
    "render_frame_rate": 24,
    "video_format": "mp4",
    "last_frame_format": "png",
    "creatomate_poll_seconds": 15,
    "creatomate_max_polls": 20,
    "times_used": 0,
}

ROWS = [
    {
        "creation_id": "PBVita-MolSmoke-01",
        "rank": 1,
        "compound_name": "GHK-Cu",
        "look_name": "Copper star (teal + copper)",
        "model_still": "openai/gpt-image-2.5-sunburst",
        "still_prompt": (
            STILL_HEADER
            + "HERO: one copper ion glowing like a miniature star, molten orange-gold with a soft "
            "corona, held at the center of a compact cradle formed by three amino acids joined in a "
            "short chain. Every atom of the peptide is a flawless sphere of clear glass that refracts "
            "the ion's light, so the whole peptide glows warm copper from the inside. The bonds "
            "between atoms are slim rods of bright light. One long, flexible side arm of the peptide "
            "curls away from the cradle into the dark. Tiny copper motes drift off the ion. The hero "
            "sits in the upper-middle of the frame, tack sharp. "
            "WORLD: a vast space between towering, rope-like collagen fibers with fine banded "
            "texture, each lit along one edge in teal and receding into deep haze. Behind them, the "
            "curved wall of a living cell glows with faint bioluminescent threads. Soft luminous "
            "vesicles float at different depths as out-of-focus spheres. "
            "LIGHT: the copper ion is the key light and throws warm caustics across the nearest "
            "fibers; a cool teal rim light from behind separates the peptide from the background; "
            "thin volumetric light shafts cut through fine haze; rich, deep blacks. "
            "CAMERA: low angle looking slightly up at the hero, shallow depth of field, creamy "
            "anamorphic oval bokeh in the foreground and background. The lower fifth of the frame "
            "stays darker and calm. "
            "COLOR: teal and copper with warm gold highlights, high contrast, filmic highlight "
            "roll-off. "
            + STILL_LOCK
        ),
        "video_motion_prompt": (
            "One continuous camera move with no cuts. The camera glides slowly forward and arcs to "
            "the left around the copper-peptide cradle at a steady, unhurried speed from the first "
            "frame to the last. At about the fifth second the three amino acids draw in closer "
            "around the copper ion: every glass atom brightens in turn, the bonds flare bright "
            "copper-white, and a soft ring of light ripples outward across the collagen fibers. The "
            "long side arm sways slowly. Copper motes drift upward, vesicles float past at different "
            "depths, and the bioluminescent threads pulse gently. The peptide keeps its shape and "
            "glass atoms, and the ion keeps its molten copper glow. The camera is still moving on "
            "the last frame."
        ),
        "extend_motion_prompt": (
            "One continuous camera move with no cuts. Continue the same slow leftward arc around the "
            "glowing copper-peptide cradle at the same steady speed, then gradually rise and pull "
            "back to reveal the full scale of the world: towering collagen fibers, the glowing wall "
            "of the cell, and vesicles drifting in haze. The copper ion pulses slowly like a "
            "heartbeat and light keeps shimmering along the bonds. Far in the distance, a few more "
            "copper points of light glint into view. The cradle keeps its exact shape, colors and "
            "glow for the whole shot. The camera is still drifting on the last frame."
        ),
    },
    {
        "creation_id": "PBVita-MolSmoke-02",
        "rank": 2,
        "compound_name": "BPC-157",
        "look_name": "Rising chain (cyan + violet + white-gold)",
        "model_still": "grok-imagine-image-2.0",
        "still_prompt": (
            STILL_HEADER
            + "HERO: a long chain of amino acids being built link by link. It rises out of a "
            "colossal ribosome at the bottom of the frame, a dark, intricately folded protein "
            "machine lit from within by electric cyan light seeping through its seams. The chain "
            "climbs through the center of the frame in a graceful, gently kinked spiral. Each amino "
            "acid is a small cluster of glossy iridescent glass spheres in violet and ice blue, and "
            "each newly formed peptide bond flashes white-gold where two amino acids meet. Loose "
            "amino acids drift in from the edges trailing thin streaks of light. The newest bond, in "
            "the upper-middle of the frame, is tack sharp. "
            "WORLD: a cavernous interior of stacked, rippling membrane sheets that glow electric "
            "cyan along their edges and recede into darkness, with soft glowing spheres hanging at "
            "different depths in fine haze. "
            "LIGHT: the white-gold bond flashes are the brightest highlights; electric cyan edge "
            "light rims every glass atom; a deep violet glow rises from below; volumetric light "
            "beams cut through the haze; rich, deep blacks. "
            "CAMERA: low angle looking up along the rising chain, shallow depth of field, "
            "anamorphic bokeh from out-of-focus spheres in the foreground. The ribosome fills the "
            "lower third in deep shadow. "
            "COLOR: electric cyan, violet and white-gold, high contrast, filmic highlight roll-off. "
            + STILL_LOCK
        ),
        "video_motion_prompt": (
            "One continuous camera move with no cuts. The camera rises slowly and steadily up along "
            "the growing amino-acid chain while drifting around it to the right, at a constant "
            "speed from the first frame to the last. New amino acids emerge from the ribosome below "
            "and lock onto the chain one after another; each new peptide bond flashes white-gold. "
            "Loose amino acids drift in from the edges on thin light streaks and join the chain. "
            "The membrane sheets ripple slowly and their cyan edges shimmer. The chain keeps its "
            "iridescent glass atoms and colors. The camera is still moving on the last frame."
        ),
        "extend_motion_prompt": (
            "One continuous camera move with no cuts. Keep rising and drifting to the right at the "
            "same steady speed. A bright wave of white-gold light travels along the full length of "
            "the chain from the ribosome to its tip, lighting every bond in turn. Then the camera "
            "slowly pulls back to reveal the whole scene: the colossal ribosome glowing cyan through "
            "its seams, the finished chain spiraling above it, and stacked membrane sheets receding "
            "into darkness. The chain keeps its shape, colors and glass atoms for the whole shot. "
            "The camera is still drifting on the last frame."
        ),
    },
]

# Lexicon: .cursor/skills/prompt-moderation-hygiene/references/lexicon.md, plus `vein`
# (injection-adjacent on a body-interior scene).
RISKY = [
    "needle", "syringe", "inject", "injection", "jab", "drug", "narcotic", "steroid",
    "blood", "gore", "wound", "serum", "plasma", "vein", "veins",
    "crash", "impact", "collide", "collision", "wreck", "debris", "explosion", "explode",
    "blast", "bomb", "fire", "flames", "burning", "shoot", "weapon", "missile", "violent",
    "crystal", "crystalline", "child", "kid", "baby", "teen", "celebrity", "wolverine",
]

GPT_SIZE_MIN_PX = 655_360
GPT_SIZE_MAX_PX = 8_294_400
GPT_SIZE_STABLE_PX = 3_686_400
GPT_EDGE_MAX = 3840


def ratio(aspect: str) -> tuple[int, int]:
    w, h = aspect.split(":")
    return int(w), int(h)


def check_rows(rows: list[dict]) -> None:
    for r in rows:
        cid = r["creation_id"]
        for k in FIELDS:
            if k in OUTPUT_COLUMNS:
                if r[k] != "":
                    raise SystemExit(f"{cid}: output column {k} must start blank")
                continue
            if r[k] == "" or r[k] is None:
                raise SystemExit(f"{cid}: {k} is empty")

        for col in PROMPT_COLUMNS:
            text = r[col].lower()
            for word in RISKY:
                if re.search(rf"\b{re.escape(word)}\b", text):
                    raise SystemExit(f"{cid}: {col} carries risky token '{word}'")

        for col in ("video_motion_prompt", "extend_motion_prompt"):
            sent = f"{SILENT_PREFIX} {r[col]}"
            if len(sent) > KLING_PROMPT_MAX:
                raise SystemExit(f"{cid}: {col} + silent prefix is {len(sent)} chars > {KLING_PROMPT_MAX}")
            if "no cuts" not in r[col] or "last frame" not in r[col]:
                raise SystemExit(f"{cid}: {col} must ask for one take and motion on the last frame")
        if len(r["negative_prompt"]) > KLING_PROMPT_MAX:
            raise SystemExit(f"{cid}: negative_prompt over {KLING_PROMPT_MAX}")

        aw, ah = ratio(r["aspect_ratio"])
        if (aw, ah) != (9, 16):
            raise SystemExit(f"{cid}: 9:16 only")
        sw, sh = (int(x) for x in r["still_size"].split("x"))
        if sw * ah != sh * aw:
            raise SystemExit(f"{cid}: still_size {r['still_size']} is not {r['aspect_ratio']}")
        if sw % 16 or sh % 16 or max(sw, sh) > GPT_EDGE_MAX:
            raise SystemExit(f"{cid}: still_size edges must be multiples of 16 and <= {GPT_EDGE_MAX}")
        if not GPT_SIZE_MIN_PX <= sw * sh <= GPT_SIZE_MAX_PX:
            raise SystemExit(f"{cid}: still_size pixel count outside GPT Image range")
        if sw * sh > GPT_SIZE_STABLE_PX:
            raise SystemExit(f"{cid}: still_size above {GPT_SIZE_STABLE_PX} px is experimental")
        if r["render_width"] * ah != r["render_height"] * aw:
            raise SystemExit(f"{cid}: render size is not {r['aspect_ratio']}")
        if min(sw, sh) < r["render_width"]:
            raise SystemExit(f"{cid}: still short side {min(sw, sh)} is softer than the {r['render_width']}px video")

        if r["resolution"] != "1080p":
            raise SystemExit(f"{cid}: resolution must be 1080p (720p is banned)")
        if not r["model_video"].startswith("fal-ai/kling-video/v3/pro/"):
            raise SystemExit(f"{cid}: model_video must be fal Kling v3 Pro (native 1080p)")
        if not (r["model_still"].startswith("openai/gpt-image-") or r["model_still"].startswith("grok-imagine-image-")):
            raise SystemExit(f"{cid}: model_still must be GPT Image or Grok Imagine Image")
        for col in ("hop1_duration_seconds", "hop2_duration_seconds"):
            if not 3 <= r[col] <= 15:
                raise SystemExit(f"{cid}: {col} outside Kling's 3-15s")
        if r["hop1_duration_seconds"] + r["hop2_duration_seconds"] != 30:
            raise SystemExit(f"{cid}: hops must add up to the 30s cut")
        if not 0 <= r["cfg_scale"] <= 1:
            raise SystemExit(f"{cid}: cfg_scale outside 0-1")
        if r["last_frame_format"] not in ("png", "jpg"):
            raise SystemExit(f"{cid}: last_frame_format must be png or jpg")
        if r["video_format"] != "mp4":
            raise SystemExit(f"{cid}: video_format must be mp4")

    leads = [r["still_prompt"][len(STILL_HEADER):len(STILL_HEADER) + 80] for r in rows]
    if len(set(leads)) != len(leads):
        raise SystemExit("rows share the same hero lead")
    for col in ("video_motion_prompt", "extend_motion_prompt"):
        if len({r[col] for r in rows}) != len(rows):
            raise SystemExit(f"rows share the same {col}")
    for r in rows:
        if r["video_motion_prompt"] == r["extend_motion_prompt"]:
            raise SystemExit(f"{r['creation_id']}: hop 2 reuses the hop 1 prompt")


def write_xlsx(rows: list[dict], path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    ws = wb.active
    ws.title = TAB
    ws.append(FIELDS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.number_format = "@"
    for r in rows:
        ws.append([r[k] if k in NUMERIC else str(r[k]) for k in FIELDS])
    wrap = set(PROMPT_COLUMNS)
    for col_idx, key in enumerate(FIELDS, start=1):
        letter = ws.cell(row=1, column=col_idx).column_letter
        ws.column_dimensions[letter].width = 60 if key in wrap else max(12, min(40, len(key) + 4))
        for row_idx in range(2, len(rows) + 2):
            cell = ws.cell(row=row_idx, column=col_idx)
            if key not in NUMERIC:
                cell.number_format = "@"
            cell.alignment = Alignment(wrap_text=key in wrap, vertical="top")
    ws.freeze_panes = "B2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", type=Path, help="also write the Drive upload copy here")
    args = ap.parse_args()

    rows = []
    for spec in ROWS:
        row = {k: "" for k in FIELDS}
        row.update(SHARED)
        row.update(spec)
        missing = set(row) - set(FIELDS)
        if missing:
            raise SystemExit(f"unknown columns {sorted(missing)}")
        rows.append(row)
    if len(rows) != 2:
        raise SystemExit("the smoke sheet holds exactly 2 rows")
    check_rows(rows)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    with OUT.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=FIELDS, lineterminator="\n")
        w.writeheader()
        w.writerows(rows)
    print(f"wrote {OUT.relative_to(ROOT.parent)}: {len(rows)} rows x {len(FIELDS)} columns")

    if args.xlsx:
        write_xlsx(rows, args.xlsx)
        print(f"wrote {args.xlsx}")

    for r in rows:
        print(
            f"  {r['creation_id']}  {r['compound_name']:<8} still={r['model_still']:<30} "
            f"still_prompt={len(r['still_prompt'])} "
            f"hop1={len(SILENT_PREFIX) + 1 + len(r['video_motion_prompt'])} "
            f"hop2={len(SILENT_PREFIX) + 1 + len(r['extend_motion_prompt'])} chars sent"
        )


if __name__ == "__main__":
    main()
